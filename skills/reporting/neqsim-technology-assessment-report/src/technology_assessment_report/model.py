"""Technology assessment report builder.

Turns a ``technology_assessment.v1`` mapping (see the technology-assessment skill) and
an optional ``technology_fmirr.v1`` mapping (see the technology-fmirr skill) into a
structured report that can be rendered as Markdown or Word (.docx).

The chapter structure follows the public technology-qualification process: the
qualification basis, technology composition, novelty categorisation and TRL, role
(enabling or enhancing), step-out, benefits and risks, standardisation, a first
failure-mode register, a preliminary qualification outline, conclusions and open points.

Markdown output needs no dependency. Word output needs ``python-docx``.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping, Sequence

REPORT_SCHEMA_VERSION = "technology_assessment_report.v1"

NOVELTY_TABLE = (
    ("Known", "1", "2", "3"),
    ("Limited knowledge", "2", "3", "4"),
    ("New", "3", "4", "4"),
)

TRL_ROWS = (
    ("0", "Unproven idea", "Paper concept; no analysis or testing."),
    ("1", "Concept demonstrated", "Basic functionality shown by analysis, analogy or sub-component tests."),
    ("2", "Concept validated", "Novel features validated by models or small-scale laboratory tests."),
    ("3", "New technology tested", "First version built and tested over a limited range of conditions."),
    ("4", "Qualified for first use", "Large-scale version tested in the intended environment; scale effects understood."),
    ("5", "Integration tested", "Full-scale unit integrated in its environment with interface and function tests."),
    ("6", "In operation", "Operated to predefined performance criteria for a limited period."),
    ("7", "Proven", "Operated to performance and reliability criteria long enough to reveal time effects."),
)

DEFAULT_REFERENCES = (
    "DNV-RP-A203, Technology qualification.",
    "ISO 20815, Production assurance and reliability management (technology qualification, "
    "TRL and novelty categorisation annex).",
    "ISO 14224, Collection and exchange of reliability and maintenance data for equipment.",
)

_CATEGORY_TEXT = {
    "1": "Proven",
    "2a": "Qualified by the user company",
    "2b": "Qualified by others",
    "2c": "Limited history, unqualified",
    "3": "New technical challenges",
    "4": "Demanding new technical challenges",
}


@dataclass
class Block:
    """One content block: paragraph, bullet list or table."""

    kind: str
    text: str = ""
    items: list[str] = field(default_factory=list)
    caption: str = ""
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    risk_column: int | None = None


@dataclass
class Section:
    """A numbered chapter of the report."""

    section_id: str
    title: str
    blocks: list[Block] = field(default_factory=list)
    landscape: bool = False

    def para(self, text: str) -> None:
        """Append a paragraph."""
        self.blocks.append(Block("p", text=text))

    def bullets(self, items: Sequence[str]) -> None:
        """Append a bullet list; empty lists are ignored."""
        if items:
            self.blocks.append(Block("bullets", items=[str(i) for i in items]))

    def table(
        self,
        caption: str,
        headers: Sequence[str],
        rows: Sequence[Sequence[Any]],
        risk_column: int | None = None,
    ) -> None:
        """Append a table."""
        self.blocks.append(
            Block(
                "table",
                caption=caption,
                headers=list(headers),
                rows=[[_cell(v) for v in row] for row in rows],
                risk_column=risk_column,
            )
        )


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return "{:g}".format(value)
    if isinstance(value, (list, tuple)):
        return "; ".join(str(v) for v in value)
    return str(value)


@dataclass
class ReportMeta:
    """Document control data."""

    title: str
    subtitle: str = ""
    document_number: str = ""
    revision: str = "A"
    author: str = ""
    reviewer: str = ""
    date: str = ""
    classification: str = "Internal"
    status: str = "Draft"
    project: str = ""


@dataclass
class Report:
    """A structured report with Markdown and Word renderers."""

    meta: ReportMeta
    sections: list[Section]
    issues: list[str]

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serialisable view of the report."""
        return {
            "schema_version": REPORT_SCHEMA_VERSION,
            "meta": vars(self.meta),
            "issues": list(self.issues),
            "sections": [
                {
                    "id": s.section_id,
                    "title": s.title,
                    "landscape": s.landscape,
                    "blocks": [vars(b) for b in s.blocks],
                }
                for s in self.sections
            ],
        }

    def to_markdown(self) -> str:
        """Render the report as Markdown."""
        m = self.meta
        out = ["# " + m.title, ""]
        if m.subtitle:
            out += ["*" + m.subtitle + "*", ""]
        control = [
            ("Document number", m.document_number),
            ("Revision", m.revision),
            ("Project", m.project),
            ("Author", m.author),
            ("Reviewer", m.reviewer),
            ("Date", m.date),
            ("Classification", m.classification),
            ("Status", m.status),
        ]
        out += ["| Item | Value |", "| --- | --- |"]
        out += ["| {} | {} |".format(k, v) for k, v in control if v]
        out.append("")
        for index, section in enumerate(self.sections, start=1):
            out += ["## {}. {}".format(index, section.title), ""]
            table_no = 0
            for block in section.blocks:
                if block.kind == "p":
                    out += [block.text, ""]
                elif block.kind == "bullets":
                    out += ["- " + item for item in block.items]
                    out.append("")
                elif block.kind == "table":
                    table_no += 1
                    out += ["**Table {}.{}: {}**".format(index, table_no, block.caption), ""]
                    out.append("| " + " | ".join(block.headers) + " |")
                    out.append("| " + " | ".join("---" for _ in block.headers) + " |")
                    for row in block.rows:
                        out.append("| " + " | ".join(c.replace("|", "/").replace("\n", " ") for c in row) + " |")
                    out.append("")
        return "\n".join(out).rstrip() + "\n"

    def save_markdown(self, path: str | Path) -> Path:
        """Write the Markdown rendering to ``path`` and return it."""
        target = Path(path)
        target.write_text(self.to_markdown(), encoding="utf-8")
        return target

    def save_docx(self, path: str | Path) -> Path:
        """Write a Word document to ``path`` and return it.

        Raises:
            ImportError: If ``python-docx`` is not installed.
        """
        return _write_docx(self, Path(path))


def _risk_counts(fmirr: Mapping[str, Any]) -> str:
    counts = fmirr["summary"]["counts"]
    return "{} High, {} Medium, {} Low".format(counts["High"], counts["Medium"], counts["Low"])


def check_report_inputs(
    assessment: Mapping[str, Any], fmirr: Mapping[str, Any] | None = None
) -> list[str]:
    """List issues that should be resolved before the report is issued.

    Args:
        assessment: A ``technology_assessment.v1`` mapping.
        fmirr: Optional ``technology_fmirr.v1`` mapping.

    Returns:
        A list of human-readable issues; empty when the inputs are consistent.

    Raises:
        ValueError: If ``assessment`` is not a ``technology_assessment.v1`` mapping.
    """
    if assessment.get("schema_version") != "technology_assessment.v1":
        raise ValueError("assessment must be a technology_assessment.v1 mapping")
    issues: list[str] = []
    issues += ["Assessment flag: " + f for f in assessment.get("flags", [])]
    issues += ["Assessment data gap: " + g for g in assessment.get("data_gaps", [])]
    novel = [e for e in assessment["elements"] if e["qualification_required"]]
    if fmirr is None:
        if novel:
            issues.append("No FMIRR supplied although {} element(s) need qualification.".format(len(novel)))
        return issues
    if fmirr.get("schema_version") != "technology_fmirr.v1":
        raise ValueError("fmirr must be a technology_fmirr.v1 mapping")
    if fmirr.get("technology") != assessment.get("technology"):
        issues.append("FMIRR technology name differs from the assessment technology name.")
    items = {str(r["item"]).strip().lower() for r in fmirr["rows"]}
    for element in novel:
        if element["name"].strip().lower() not in items and element["id"].strip().lower() not in items:
            issues.append("Novel element {} ({}) has no failure mode in the FMIRR.".format(
                element["id"], element["name"]))
    for finding in fmirr.get("findings", []):
        if finding["level"] in ("gap", "warning"):
            issues.append("FMIRR {}: {} {}".format(finding["level"], finding["entry_id"], finding["text"]).strip())
    return issues


def _executive_summary(
    assessment: Mapping[str, Any], fmirr: Mapping[str, Any] | None, sec: Section
) -> None:
    s = assessment["summary"]
    sec.para(
        "This report assesses the technology \"{}\". The technology was decomposed into {} element(s), "
        "of which {} require qualification. The highest novelty category is {} and the system technology "
        "readiness level (TRL) is {} ({}).".format(
            assessment["technology"], s["element_count"], s["novel_element_count"],
            s["highest_category"], s["system_trl"], s["system_trl_name"])
    )
    sec.para("The technology is classified as {}. {}".format(s["role"], s["role_rationale"]))
    beyond = [r["dimension"] for r in assessment["step_out"] if r["status"] == "step-out beyond experience"]
    if beyond:
        sec.para("Step-out beyond existing experience was identified for: {}.".format(", ".join(beyond)))
    if fmirr is not None:
        fs = fmirr["summary"]
        sec.para(
            "The first failure-mode identification and risk ranking registered {} failure mode(s): {}. "
            "{} critical mode(s) need qualification activities; {} of them are linked to a test.".format(
                fs["total"], _risk_counts(fmirr), fs["critical_count"], fs["critical_with_test"])
        )
    sec.para(
        "The conclusions are screening-level and depend on the evidence available. Open points are listed "
        "in the final chapters."
    )


def _decomposition_rows(elements: Sequence[Mapping[str, Any]]) -> list[list[Any]]:
    children: dict[Any, list[Mapping[str, Any]]] = {}
    for e in elements:
        children.setdefault(e["parent_id"], []).append(e)
    rows: list[list[Any]] = []

    def walk(parent: Any, depth: int) -> None:
        for e in children.get(parent, []):
            rows.append([e["id"], "\u00a0\u00a0" * depth + e["name"], e["kind"], e["function"]])
            walk(e["id"], depth + 1)

    walk(None, 0)
    return rows


def build_report(
    assessment: Mapping[str, Any],
    fmirr: Mapping[str, Any] | None = None,
    meta: ReportMeta | None = None,
    *,
    narrative: Mapping[str, Sequence[str]] | None = None,
    references: Sequence[str] = (),
    introduction_context: str = "",
) -> Report:
    """Build the technology assessment report.

    Args:
        assessment: A ``technology_assessment.v1`` mapping.
        fmirr: Optional ``technology_fmirr.v1`` mapping.
        meta: Document control data; a default title is generated when omitted.
        narrative: Optional author-written paragraphs per section id (for example
            ``{"introduction": [...], "conclusions": [...]}``), appended after the generated text.
        references: Extra references appended to the defaults.
        introduction_context: One paragraph on the project and why the technology is considered.

    Returns:
        A :class:`Report` ready to render.

    Raises:
        ValueError: If the inputs are not the expected schema versions.
    """
    issues = check_report_inputs(assessment, fmirr)
    name = assessment["technology"]
    meta = meta or ReportMeta(title="Technology Assessment Report: " + name)
    if not meta.date:
        meta.date = _dt.date.today().isoformat()
    narrative = narrative or {}
    basis = assessment["qualification_basis"]
    summary = assessment["summary"]
    elements = assessment["elements"]
    sections: list[Section] = []

    def new(section_id: str, title: str, landscape: bool = False) -> Section:
        sec = Section(section_id, title, landscape=landscape)
        sections.append(sec)
        return sec

    sec = new("executive_summary", "Executive summary")
    _executive_summary(assessment, fmirr, sec)

    sec = new("introduction", "Introduction")
    if introduction_context:
        sec.para(introduction_context)
    sec.para(
        "The purpose of the technology assessment is to determine which elements of the technology "
        "require qualification and to identify their main challenges and uncertainties. The assessment "
        "covers the technology qualification basis, technology composition, novelty categorisation, "
        "technology readiness level, benefits and risks, step-out, standardisation potential and a "
        "preliminary outline of qualification activities."
    )
    sec.para(
        "The method follows DNV-RP-A203 and ISO 20815. Novelty and TRL ratings are expert judgement and "
        "count only documented, traceable and accessible evidence."
    )

    sec = new("qualification_basis", "Technology qualification basis")
    sec.para(
        "The qualification basis is the common ground and set of criteria against which the qualification "
        "activities and decisions are assessed. It is updated as information on the technology becomes available."
    )
    sec.table(
        "Technology qualification basis",
        ["Topic", "Content"],
        [
            ["Description of the technology", basis["description"] or "Not stated"],
            ["How the technology is used", basis["intended_use"] or "Not stated"],
            ["Intended environment", basis["environment"] or "Not stated"],
            ["Functions", basis["functions"] or "Not stated"],
            ["Performance expectations", basis["performance_expectations"] or "Not stated"],
            ["Regulatory requirements", basis["regulatory_requirements"] or "Not stated"],
        ],
    )

    sec = new("composition", "Technology description and decomposition")
    sec.para("The technology is decomposed top-down from the system function to sub-systems, components and software.")
    sec.table("Technology decomposition", ["Id", "Element", "Kind", "Function"], _decomposition_rows(elements))

    sec = new("maturity", "Technology maturity: novelty category and TRL")
    sec.para(
        "Each element is categorised by the novelty of the technology and its application area. Categories "
        "2 to 4 require qualification; category 1 elements are verified to the applicable standard."
    )
    sec.table(
        "Novelty and TRL per element",
        ["Id", "Element", "Technology novelty", "Application area", "Category", "TRL", "Qualification"],
        [
            [
                e["id"], e["name"], e["technology_novelty"].replace("_", " "),
                e["application_area"].replace("_", " "),
                "{} ({})".format(e["category_label"], _CATEGORY_TEXT[e["category_label"]]),
                "{} {}".format(e["trl"], e["trl_name"]),
                "Required ({} risk)".format(e["qualification_risk"]) if e["qualification_required"] else "Not required",
            ]
            for e in elements
        ],
    )
    sec.para("System TRL: {} ({}). Basis: {}".format(summary["system_trl"], summary["system_trl_name"], summary["system_trl_basis"]))
    flagged = [(e["id"], f) for e in elements for f in e["flags"]]
    if flagged:
        sec.para("Consistency flags raised on the ratings:")
        sec.bullets(["{}: {}".format(i, f) for i, f in flagged])

    sec = new("role", "Enabling or enhancing technology")
    sec.para("Classification: {}. {}".format(summary["role"], summary["role_rationale"]))

    sec = new("step_out", "Step-out from existing experience")
    sec.table(
        "Step-out per dimension",
        ["Dimension", "Experience basis", "Proposed", "Reference", "Proposed value", "Status"],
        [
            [
                r["dimension"].replace("_", " "), r["experience_basis"], r["proposed"],
                "{} {}".format(_cell(r["reference_value"]), r["unit"]).strip(),
                "{} {}".format(_cell(r["proposed_value"]), r["unit"]).strip(), r["status"],
            ]
            for r in assessment["step_out"]
        ]
        or [["No step-out recorded", "", "", "", "", ""]],
    )

    sec = new("benefits_risks", "Benefits and risks")
    sec.para("Possible benefits of using the technology:")
    sec.bullets([b["text"] + (" ({})".format(b["kind"]) if b["kind"] else "") for b in assessment["benefits"]] or ["None recorded."])
    sec.para("Possible risks of using the technology:")
    sec.table(
        "Risks and mitigations",
        ["Risk", "Category", "Mitigation"],
        [[r["text"], r["kind"], r["mitigation"]] for r in assessment["risks"]] or [["None recorded", "", ""]],
    )

    sec = new("standardisation", "Standardisation and multi-implementation")
    sec.bullets(assessment["standardisation"] + assessment["multi_implementation"] or ["Not assessed."])

    if fmirr is not None:
        sec = new("fmirr", "Failure mode identification and risk ranking (first version)", landscape=True)
        fs = fmirr["summary"]
        sec.para(
            "Failure modes are registered with failed item, failure mode, mechanism and cause kept separate. "
            "Result: {}. Critical (Medium and High) modes: {}.".format(_risk_counts(fmirr), ", ".join(fs["critical_ids"]) or "none")
        )
        sec.table(
            "FMIRR register",
            ["Id", "Item", "Function", "Failure mode", "Mechanism", "Cause", "Detection", "Consequence",
             "Cons.", "Prob.", "Risk", "Test"],
            [
                [r["id"], r["item"], r["function"], r["failure_mode"], r["failure_mechanism"], r["failure_cause"],
                 r["detection"], r["consequence"], r["consequence_class"], r["probability_class"], r["risk"], r["test"]]
                for r in fmirr["rows"]
            ],
            risk_column=10,
        )
        grid = fmirr["risk_matrix"]["rows_probability_5_to_1"]
        sec.table(
            "Number of failure modes per probability and consequence class",
            ["Probability \\ Consequence", "1", "2", "3", "4", "5"],
            [[p] + counts for p, counts in grid.items()],
        )

    sec = new("qualification_outline", "Preliminary outline of qualification activities")
    outline = assessment["qualification_outline"]
    if outline:
        sec.para("The activities below target TRL 4 (qualified for first use). TRL 5 to 7 follow by application.")
        sec.table(
            "Qualification outline per element",
            ["Element", "Category", "Current TRL", "Target TRL", "Activities"],
            [[o["element"], o["category"], o["current_trl"], o["target_trl"], o["activities"]] for o in outline],
        )
        if fmirr is not None:
            linked = [r for r in fmirr["rows"] if r["critical"] and r["test"]]
            if linked:
                sec.para("Critical failure modes with a linked qualification activity:")
                sec.bullets(["{}: {} -> {}".format(r["id"], r["failure_mode"], r["test"]) for r in linked])
    else:
        sec.para("No element requires technology qualification; elements are verified to the applicable standards.")

    sec = new("conclusions", "Conclusions and recommendations")
    novel = sorted((e for e in elements if e["qualification_required"]), key=lambda e: (-e["category"], e["trl"]))
    if novel:
        sec.para("Elements that need qualification, ordered by novelty and maturity:")
        sec.bullets(["{} {}: category {}, TRL {}".format(e["id"], e["name"], e["category_label"], e["trl"]) for e in novel])
        if summary["role"] == "enabling":
            sec.para(
                "The technology is enabling: the critical path of the project may run through the qualification. "
                "Plan the qualification programme before the concept is locked and keep a fallback open where possible."
            )
        sec.para("Recommended next steps: confirm the ratings in an expert-panel workshop, repeat the FMIRR with "
                 "the panel, and prepare the technology qualification programme.")
    else:
        sec.para("No qualification is required. Verify all elements to the applicable standards and practice.")
    sec.para("Category 1 elements are not exempt from verification: verify them to the applicable standard with ordinary design review and QA/QC.")

    sec = new("open_points", "Assumptions, data gaps and open points")
    sec.para("Assumptions:")
    sec.bullets(assessment["assumptions"] + (fmirr["assumptions"] if fmirr is not None else []))
    sec.para("Data gaps and open issues to resolve before issue:")
    sec.bullets(issues or ["None."])

    sec = new("references", "References")
    sec.bullets(list(DEFAULT_REFERENCES) + list(references))

    sec = new("appendix_definitions", "Appendix A: Definitions")
    sec.table(
        "Technology novelty categorisation",
        ["Application area", "Proven", "Limited history", "New or unproven"],
        [list(r) for r in NOVELTY_TABLE],
    )
    sec.table("Technology readiness levels", ["TRL", "Stage", "Description"], [list(r) for r in TRL_ROWS])
    if fmirr is not None:
        sec.table("Consequence classes", ["Class", "Description"], list(fmirr["scales"]["consequence"].items()))
        sec.table("Probability classes", ["Class", "Description"], list(fmirr["scales"]["probability"].items()))

    for section in sections:
        extra = narrative.get(section.section_id)
        if extra:
            for paragraph in extra:
                section.para(paragraph)
    return Report(meta=meta, sections=sections, issues=issues)


_RISK_FILL = {"Low": "C6EFCE", "Medium": "FFEB9C", "High": "FFC7CE"}


def _column_widths(
    headers: Sequence[str], rows: Sequence[Sequence[str]], usable_inches: float, font_pt: float
) -> list[float]:
    """Return column widths in inches that fit ``usable_inches``.

    Every column gets enough width for its longest word (capped) plus cell padding; the remaining
    width is shared in proportion to the average text length of each column.

    Args:
        headers: Column headers.
        rows: Table rows as strings.
        usable_inches: Width available between the page margins.
        font_pt: Font size in points used in the table.

    Returns:
        One width per column; the widths sum to ``usable_inches``.
    """
    char = font_pt * 0.0085
    minimums, flex = [], []
    for index, header in enumerate(headers):
        cells = [row[index] for row in rows if index < len(row)]
        words = [len(word) for text in cells + [header] for word in text.split()] or [3]
        minimums.append(min(max(words), 14) * char + 0.2)
        average = sum(len(c) for c in cells) / len(cells) if cells else 0.0
        flex.append(max(min(average, 40), 1.0))
    total_min = sum(minimums)
    if total_min >= usable_inches:
        return [m * usable_inches / total_min for m in minimums]
    spare = usable_inches - total_min
    return [m + spare * f / sum(flex) for m, f in zip(minimums, flex)]


def _table_font_size(columns: int) -> int | None:
    """Return a table font size in points for a table with the given number of columns."""
    if columns >= 9:
        return 8
    if columns >= 5:
        return 9
    return None


def _write_docx(report: Report, path: Path) -> Path:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT, WD_SECTION
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Emu, Inches, Pt
    except ImportError as exc:
        raise ImportError("python-docx is required for Word output: pip install python-docx") from exc

    def field_run(paragraph: Any, code: str) -> None:
        run = paragraph.add_run()
        for kind, text in (("begin", None), (None, code), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    def shade(cell: Any, fill: str) -> None:
        props = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        props.append(shd)

    def repeat_header(row: Any) -> None:
        props = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        props.append(el)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    meta = report.meta

    doc.add_heading(meta.title, 0)
    if meta.subtitle:
        doc.add_paragraph(meta.subtitle)
    control = [
        ("Document number", meta.document_number), ("Revision", meta.revision), ("Project", meta.project),
        ("Author", meta.author), ("Reviewer", meta.reviewer), ("Date", meta.date),
        ("Classification", meta.classification), ("Status", meta.status),
    ]
    control = [(k, v) for k, v in control if v]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in control:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    toc_heading = doc.add_paragraph()
    toc_heading.add_run("Contents").bold = True
    field_run(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_page_break()

    first = doc.sections[0]
    first.header.paragraphs[0].text = meta.title
    footer = first.footer.paragraphs[0]
    footer.text = "Security classification: {} - Status: {} - Page ".format(meta.classification, meta.status)
    field_run(footer, "PAGE")
    footer.add_run(" of ")
    field_run(footer, "NUMPAGES")

    landscape_open = False
    for index, section in enumerate(report.sections, start=1):
        if section.landscape and not landscape_open:
            new_section = doc.add_section(WD_SECTION.NEW_PAGE)
            new_section.orientation = WD_ORIENT.LANDSCAPE
            new_section.page_width, new_section.page_height = new_section.page_height, new_section.page_width
            new_section.left_margin = new_section.right_margin = Inches(0.7)
            landscape_open = True
        elif not section.landscape and landscape_open:
            new_section = doc.add_section(WD_SECTION.NEW_PAGE)
            new_section.orientation = WD_ORIENT.PORTRAIT
            new_section.page_width, new_section.page_height = new_section.page_height, new_section.page_width
            new_section.left_margin = new_section.right_margin = doc.sections[0].left_margin
            landscape_open = False
        doc.add_heading("{}. {}".format(index, section.title), 1)
        table_no = 0
        for block in section.blocks:
            if block.kind == "p":
                doc.add_paragraph(block.text)
            elif block.kind == "bullets":
                for item in block.items:
                    doc.add_paragraph(item, style="List Bullet")
            elif block.kind == "table":
                table_no += 1
                caption = doc.add_paragraph()
                caption.add_run("Table {}.{}: {}".format(index, table_no, block.caption)).bold = True
                tbl = doc.add_table(rows=1, cols=len(block.headers))
                tbl.style = "Table Grid"
                tbl.autofit = False
                for cell, header in zip(tbl.rows[0].cells, block.headers):
                    cell.text = header
                    cell.paragraphs[0].runs[0].bold = True
                    shade(cell, "D9D9D9")
                repeat_header(tbl.rows[0])
                for row in block.rows:
                    cells = tbl.add_row().cells
                    for position, (cell, value) in enumerate(zip(cells, row)):
                        cell.text = value
                        if block.risk_column == position and value in _RISK_FILL:
                            shade(cell, _RISK_FILL[value])
                current = doc.sections[-1]
                usable = current.page_width - current.left_margin - current.right_margin
                size = _table_font_size(len(block.headers))
                widths = _column_widths(block.headers, block.rows, usable / 914400.0, size or 10)
                for position, inches in enumerate(widths):
                    width = Emu(int(inches * 914400))
                    tbl.columns[position].width = width
                    for row in tbl.rows:
                        row.cells[position].width = width
                if size is not None:
                    for row in tbl.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.size = Pt(size)
                doc.add_paragraph()
    doc.save(str(path))
    return path
