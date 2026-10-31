"""Word report builder for a well design study (python-docx). Content is supplied as sections; numbering is automatic."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


STATUS_COLOURS = {
    "[VERIFIED]": RGBColor(0x1B, 0x7F, 0x3B),
    "[UNVERIFIED-SOURCE]": RGBColor(0xB3, 0x6B, 0x00),
    "[DERIVED]": RGBColor(0x1F, 0x4E, 0x9E),
    "[ASSUMED]": RGBColor(0xB0, 0x2A, 0x2A),
    "[PLACEHOLDER]": RGBColor(0x66, 0x66, 0x66),
    "[NEEDS_REVIEW]": RGBColor(0xB0, 0x2A, 0x2A),
}


def _field(par, code: str, placeholder: str = "1") -> None:
    r = par.add_run()
    for kind, txt in (("begin", None), ("instr", code), ("separate", None), ("text", placeholder), ("end", None)):
        if kind in ("begin", "separate", "end"):
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        elif kind == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        else:
            e = OxmlElement("w:t")
            e.text = txt
        r._r.append(e)


class WellReport:
    """Collects headings, paragraphs, tables and figures and writes a Word document with contents and page numbers."""

    def __init__(self, title: str, subtitle: str, doc_id: str, status: str) -> None:
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, a, Cm(1.9))
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{doc_id} - {status} - page ")
        r.font.size = Pt(8)
        _field(p, "PAGE")
        self.ch = self.h2 = 0
        self.app_n = 0
        self.usable_cm = 17.2
        self.fig_n = 0
        self.tab_n = 0
        self.figures: List[Dict[str, str]] = []
        for _ in range(5):
            self.doc.add_paragraph()
        t = self.doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = t.add_run(title)
        rr.bold = True
        rr.font.size = Pt(26)
        t = self.doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.add_run(subtitle).font.size = Pt(13)
        for line in (doc_id, status):
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.add_run(line).font.size = Pt(11)
        self.doc.add_page_break()
        cp = self.doc.add_paragraph()
        cr = cp.add_run("Contents")
        cr.bold = True
        cr.font.size = Pt(16)
        tp = self.doc.add_paragraph()
        _field(tp, 'TOC \\o "1-2" \\h \\z \\u', "Update the field to build the table of contents.")
        self.doc.add_page_break()

    def _runs(self, par, text: str, size: Optional[float] = None) -> None:
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[(?:VERIFIED|UNVERIFIED-SOURCE|DERIVED|ASSUMED|PLACEHOLDER|NEEDS_REVIEW)\])", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                r = par.add_run(part[2:-2])
                r.bold = True
            elif part in STATUS_COLOURS:
                r = par.add_run(part)
                r.bold = True
                r.font.color.rgb = STATUS_COLOURS[part]
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                r = par.add_run(part[1:-1])
                r.font.name = "Consolas"
            else:
                r = par.add_run(part)
            if size:
                r.font.size = Pt(size)

    def h1(self, text: str, page_break: bool = True) -> None:
        self.ch += 1
        self.h2 = 0
        if page_break and self.ch > 1:
            self.doc.add_page_break()
        self.doc.add_heading(f"{self.ch} {text}", 1)

    def h2_(self, text: str) -> None:
        self.h2 += 1
        self.doc.add_heading(f"{self.ch}.{self.h2} {text}", 2)

    def appendix(self, text: str, page_break: bool = True) -> None:
        """Start an appendix (A, B, ...) as a level-1 heading so that it appears in the contents."""
        self.app_n += 1
        if page_break:
            self.doc.add_page_break()
        self.doc.add_heading(f"Appendix {chr(64 + self.app_n)} {text}", 1)

    def landscape(self) -> None:
        """Start a new landscape section (wide tables); table widths adapt to the usable width."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        self.usable_cm = 25.9

    def portrait(self) -> None:
        """Return to a portrait section."""
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        self.usable_cm = 17.2

    def callout(self, title: str, lines: Sequence[str], fill: str = "EAF1FB", size: float = 9.5) -> None:
        """Shaded single-cell box (answer-first summary, key warnings). Lines support the same inline markup as paragraphs."""
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell.width = Cm(self.usable_cm)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(size + 1)
        for ln in lines:
            q = cell.add_paragraph()
            q.paragraph_format.space_after = Pt(2)
            self._runs(q, ln, size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def para(self, text: str, size: Optional[float] = None) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        self._runs(p, text, size)

    def bullets(self, items: Sequence[str], size: Optional[float] = None) -> None:
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            self._runs(p, it, size)

    def numbered(self, items: Sequence[str]) -> None:
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(f"{i}. ")
            self._runs(p, it)

    def table(self, rows: Sequence[Sequence[object]], caption: Optional[str] = None, font: float = 8.0, widths: Optional[Sequence[float]] = None) -> None:
        rows = [[("" if c is None else str(c)) for c in r] for r in rows]
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        if caption:
            self.tab_n += 1
            cp = self.doc.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            rr = cp.add_run(f"Table {self.tab_n}. {caption}")
            rr.bold = True
            rr.font.size = Pt(9)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        usable = self.usable_cm
        if widths is None:
            lens = [min(max(max(len(r[c]) for r in rows), 6), 60) ** 0.8 for c in range(ncol)]
            tot = sum(lens)
            widths = [usable * x / tot for x in lens]
        for i, row in enumerate(rows):
            tr = t.rows[i]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for j, txt in enumerate(row):
                cell = tr.cells[j]
                cell.width = Cm(widths[j])
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                self._runs(par, txt, font)
                if i == 0:
                    for r in par.runs:
                        r.bold = True
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "D9E2F3")
                    cell._tc.get_or_add_tcPr().append(shd)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, path: Path, caption: str, provenance: str = "", width_cm: float = 15.5) -> None:
        self.fig_n += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(1)
        r = c.add_run(f"Figure {self.fig_n}. {caption}")
        r.bold = True
        r.font.size = Pt(9)
        if provenance:
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = q.add_run(provenance)
            rr.italic = True
            rr.font.size = Pt(7)
            rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        self.figures.append({"n": str(self.fig_n), "file": Path(path).name, "caption": caption})

    def save(self, path: Path) -> Path:
        path = Path(path)
        self.doc.save(path)
        return path
